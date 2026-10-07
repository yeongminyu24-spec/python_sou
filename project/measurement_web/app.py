"""엑셀의 SPL 측정값을 항목별로 직접 검토하는 Flask 실습 앱."""
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from flask import Flask, render_template, request
from openpyxl import load_workbook
from werkzeug.exceptions import RequestEntityTooLarge
import re

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024  # 업로드 최대 20MB


def is_blank(value):
    # 숫자 0은 빈칸이 아닙니다.
    return value is None or (isinstance(value, str) and not value.strip())


def number(value):
    """숫자로 바꾸되 빈칸, 논리값, 무한대 등은 허용하지 않습니다."""
    if is_blank(value) or isinstance(value, bool):
        raise ValueError('숫자가 아닙니다')
    try:
        result = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError('숫자가 아닙니다') from None
    if not result.is_finite():
        raise ValueError('유한한 숫자가 아닙니다')
    return result


def display(value):
    if isinstance(value, Decimal):
        return format(value, 'f')
    return '' if value is None else str(value)


def analyze_workbook(stream, deviation_percent=Decimal('10')):
    # 수식 자체도 읽습니다. SPL/규격 수식은 캐시값을 믿지 않고 오류로 표시합니다.
    workbook = load_workbook(stream, data_only=False, read_only=True)
    items, pages, warnings = [], [], []
    try:
        for sheet in workbook.worksheets:
            pages.append(sheet.title)
            records = sheet.iter_rows(min_row=13, values_only=True)
            headers = next(records, ())
            # Q~S로 고정하지 않고 13행의 Spl 1, Spl 2 등의 이름으로 찾습니다.
            spl_columns = []
            for index, heading in enumerate(headers):
                if re.fullmatch(r'Spl\s*\d+', str(heading).strip(), re.IGNORECASE):
                    spl_columns.append((index, str(heading).strip()))
            if not spl_columns:
                warnings.append(f'{sheet.title}: 13행에서 SPL 머리글을 찾지 못했습니다.')

            for excel_row, values in enumerate(records, start=14):
                def cell(index):
                    return values[index] if index < len(values) else None

                # C~F가 모두 비어 있는 행은 빈 행/통계 행으로 보고 제외합니다.
                if all(is_blank(cell(i)) for i in range(2, 6)):
                    continue
                issues = []
                lower = upper = None
                try:
                    nominal, plus, minus = (number(cell(i)) for i in (5, 6, 7))
                    if plus < 0 or minus < 0:
                        raise ValueError('공차의 크기는 0 이상이어야 합니다')
                    lower, upper = nominal - minus, nominal + plus
                except ValueError:
                    issues.append('규격 오류')

                samples, numeric = [], []
                missing = invalid = outside = 0
                for index, label in spl_columns:
                    raw = cell(index)
                    state = '미측정'
                    if is_blank(raw):
                        missing += 1
                    else:
                        try:
                            value = number(raw)
                            numeric.append(value)
                            state = '확인 필요' if lower is None else 'OK'
                            if lower is not None and not lower <= value <= upper:
                                outside += 1
                                state = 'NG'
                        except ValueError:
                            invalid += 1
                            state = '값 오류'
                    samples.append({'label': label, 'value': display(raw), 'state': state})
                if invalid:
                    issues.append('값 오류')
                if not spl_columns:
                    issues.append('SPL 열 없음')
                elif missing == len(spl_columns):
                    issues.append('미측정')
                elif missing:
                    issues.append('일부 미측정')

                # 한 값이라도 범위 밖이면 NG를 우선합니다. 다른 문제도 함께 남깁니다.
                if outside:
                    status = 'NG'
                elif lower is None or invalid or not spl_columns:
                    status = '확인 필요'
                elif missing == len(spl_columns):
                    status = '미측정'
                elif missing:
                    status = '확인 필요'
                else:
                    status = 'OK'
                # 평균은 이 항목의 유효한 SPL에 대해서만 계산합니다.
                average = sum(numeric, Decimal(0)) / len(numeric) if numeric else None
                # 주의 표시는 규격 판정과 별개입니다. 반올림 전 평균으로 비교합니다.
                # 유효 값이 2개 이상일 때만 같은 항목 안에서 편차를 확인합니다.
                for sample in samples:
                    sample['review'] = False
                    if len(numeric) >= 2 and sample['state'] not in ('미측정', '값 오류'):
                        distance = abs(number(sample['value']) - average)
                        threshold = abs(average) * deviation_percent / Decimal(100)
                        sample['review'] = distance > threshold
                review = any(sample['review'] for sample in samples)
                items.append({
                    'page': sheet.title, 'row': excel_row,
                    'item': display(cell(2)), 'location': display(cell(3)),
                    'kind': display(cell(4)), 'lower': display(lower),
                    'upper': display(upper), 'samples': samples,
                    'average': format(average.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP), 'f') if average is not None else '',
                    'review': review, 'count': len(numeric),
                    'status': status, 'issues': ' / '.join(issues),
                    'notes': display(cell(19)),
                })
    finally:
        workbook.close()
    return items, pages, warnings


@app.route('/', methods=['GET', 'POST'])
def index():
    items, pages, warnings = [], [], []
    error = filename = ''
    deviation_percent = Decimal('10')
    if request.method == 'POST':
        try:
            deviation_percent = number(request.form.get('deviation_percent', '10'))
            if not 0 < deviation_percent <= 100:
                raise ValueError('범위 오류')
        except ValueError:
            error = '평균 편차 기준은 0 초과 100 이하의 숫자로 입력하세요.'
        uploaded = request.files.get('file')
        if error:
            pass
        elif not uploaded or not uploaded.filename:
            error = '분석할 .xlsx 파일을 선택하세요.'
        elif not uploaded.filename.lower().endswith('.xlsx'):
            error = '.xlsx 파일만 업로드할 수 있습니다.'
        else:
            filename = uploaded.filename
            try:
                items, pages, warnings = analyze_workbook(BytesIO(uploaded.read()), deviation_percent)
            except Exception:
                # 손상된 파일의 내부 오류 대신 학습자가 이해할 수 있는 안내를 표시합니다.
                error = '엑셀을 읽지 못했습니다. 암호가 없고 정상적인 .xlsx 파일인지 확인하세요.'
    summary = {key: sum(item['status'] == key for item in items)
               for key in ('OK', 'NG', '미측정', '확인 필요')}
    return render_template('index.html', items=items, pages=pages,
                           warnings=warnings, error=error, filename=filename,
                           summary=summary, deviation_percent=display(deviation_percent))


@app.errorhandler(RequestEntityTooLarge)
def too_large(error):
    return render_template('index.html', items=[], pages=[], warnings=[],
                           error='파일 크기는 20MB 이하여야 합니다.', filename='',
                           summary={key: 0 for key in ('OK', 'NG', '미측정', '확인 필요')}), 413


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=True)
